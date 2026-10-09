"""Genera Excel organizado para el plan de transferencia a partir de catalog.csv.

Si existen pipeline_schedules.csv y/o pipeline_runs.csv en la misma carpeta,
agrega las hojas "Programaciones" y "Ejecuciones".
"""
import glob
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob("output/*/catalog.csv"))[-1]
src_dir = Path(src).parent

d = pd.read_csv(src, encoding="utf-8-sig").fillna("")
d[["Proyecto", "Ambiente"]] = d.workspace.str.split("-", n=1, expand=True)
ENVS = ["PROD", "STAGE", "TEST"]
d["Ambiente"] = pd.Categorical(d["Ambiente"], ENVS, ordered=True)

inv = d.rename(columns={"workspace": "Workspace", "itemType": "Tipo", "itemName": "Item",
    "description": "Descripcion", "itemId": "ItemId", "workspaceId": "WorkspaceId"})
inv = inv.sort_values(["Proyecto", "Ambiente", "Tipo", "Item"])
inv_out = inv[["Proyecto", "Ambiente", "Workspace", "Tipo", "Item", "Descripcion", "ItemId", "WorkspaceId"]].copy()
inv_out["Responsable"] = ""; inv_out["Prioridad"] = ""; inv_out["Estado transferencia"] = ""; inv_out["Observaciones"] = ""

res_ws = d.groupby(["Proyecto", "Ambiente"], observed=True).size().unstack(fill_value=0)[ENVS]
res_ws["Total"] = res_ws.sum(axis=1)
res_ws = res_ws.reset_index()

tipo = d.pivot_table(index="itemType", columns="workspace", values="itemId", aggfunc="count", fill_value=0)
tipo["Total"] = tipo.sum(axis=1)
tipo = tipo.sort_values("Total", ascending=False).reset_index().rename(columns={"itemType": "Tipo"})

pv = d.pivot_table(index=["Proyecto", "itemType", "itemName"], columns="Ambiente", values="itemId",
                   aggfunc="count", observed=True).reindex(columns=ENVS)
par = pv.notna().replace({True: "X", False: ""})
par["Ambientes"] = pv.notna().sum(axis=1)
par = par.reset_index().rename(columns={"itemType": "Tipo", "itemName": "Item"})
par["Observacion"] = par["Ambientes"].map({3: "En los 3 ambientes", 2: "Falta en un ambiente", 1: "Solo en un ambiente"})
par = par.sort_values(["Proyecto", "Ambientes", "Tipo", "Item"], ascending=[True, True, True, True])

sql = inv_out[inv_out.Tipo.isin(["Warehouse", "Lakehouse", "SQLEndpoint"])]
orq = inv_out[inv_out.Tipo.isin(["DataPipeline", "Dataflow", "CopyJob", "Notebook"])]
bi = inv_out[inv_out.Tipo.isin(["SemanticModel", "Report", "PaginatedReport"])]

SCHED_COLS = [
    "workspaceName", "pipelineName", "scheduleId", "scheduleEnabled",
    "scheduleStatus", "scheduleType", "scheduleDescription",
    "startDateTime", "endDateTime", "localTimeZoneId",
    "convertedTime", "convertedTz", "conversionOk",
    "ownerId", "ownerType", "createdDateTime", "collectedAtUtc",
]
SCHED_RENAME = {
    "workspaceName": "Workspace", "pipelineName": "Pipeline",
    "scheduleId": "ScheduleId", "scheduleEnabled": "Habilitada",
    "scheduleStatus": "Estado", "scheduleType": "Tipo",
    "scheduleDescription": "Descripcion",
    "startDateTime": "Inicio", "endDateTime": "Fin",
    "localTimeZoneId": "TZ Original", "convertedTime": "Hora Convertida",
    "convertedTz": "TZ Conversion", "conversionOk": "Conversion OK",
    "ownerId": "OwnerId", "ownerType": "OwnerType",
    "createdDateTime": "Creada", "collectedAtUtc": "Recolectada UTC",
}

RUNS_COLS = [
    "workspaceName", "pipelineName", "runId", "invokeType", "status",
    "startTimeUtc", "endTimeUtc", "duration",
    "failureReason", "rootActivityId", "jobType", "collectedAtUtc",
]
RUNS_RENAME = {
    "workspaceName": "Workspace", "pipelineName": "Pipeline",
    "runId": "RunId", "invokeType": "Invocacion", "status": "Estado",
    "startTimeUtc": "Inicio UTC", "endTimeUtc": "Fin UTC",
    "duration": "Duracion", "failureReason": "Error",
    "rootActivityId": "ActivityId", "jobType": "JobType",
    "collectedAtUtc": "Recolectada UTC",
}

PIPE_COLS = [
    "workspaceName", "pipelineName", "pipelineId",
    "scheduleCount", "scheduleStatus",
]
PIPE_RENAME = {
    "workspaceName": "Workspace", "pipelineName": "Pipeline",
    "pipelineId": "PipelineId", "scheduleCount": "Programaciones",
    "scheduleStatus": "Estado Schedules",
}

sched_df = None
sched_path = src_dir / "pipeline_schedules.csv"
if sched_path.exists():
    sched_df = pd.read_csv(sched_path, encoding="utf-8-sig").fillna("")
    cols = [c for c in SCHED_COLS if c in sched_df.columns]
    sched_df = sched_df[cols].rename(columns=SCHED_RENAME)
    sched_df = sched_df.sort_values(["Workspace", "Pipeline", "Tipo"])

runs_df = None
runs_path = src_dir / "pipeline_runs.csv"
if runs_path.exists():
    runs_df = pd.read_csv(runs_path, encoding="utf-8-sig").fillna("")
    cols = [c for c in RUNS_COLS if c in runs_df.columns]
    runs_df = runs_df[cols].rename(columns=RUNS_RENAME)
    runs_df = runs_df.sort_values(["Workspace", "Pipeline", "Inicio UTC"], ascending=[True, True, False])

pipes_df = None
pipes_path = src_dir / "pipelines.csv"
if pipes_path.exists():
    pipes_df = pd.read_csv(pipes_path, encoding="utf-8-sig").fillna("")
    cols = [c for c in PIPE_COLS if c in pipes_df.columns]
    pipes_df = pipes_df[cols].rename(columns=PIPE_RENAME)
    pipes_df = pipes_df.sort_values(["Workspace", "Pipeline"])

out = Path("output") / "Inventario_Fabric_Transferencia.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    res_ws.to_excel(xw, sheet_name="Resumen", index=False, startrow=3)
    tipo.to_excel(xw, sheet_name="Por tipo", index=False)
    par.to_excel(xw, sheet_name="Paridad ambientes", index=False)
    inv_out.to_excel(xw, sheet_name="Inventario completo", index=False)
    orq.to_excel(xw, sheet_name="Orquestacion y ETL", index=False)
    sql.to_excel(xw, sheet_name="Datos (WH-LH)", index=False)
    bi.to_excel(xw, sheet_name="BI (Modelos-Reportes)", index=False)

    if pipes_df is not None and not pipes_df.empty:
        pipes_df.to_excel(xw, sheet_name="Pipelines", index=False)
    if sched_df is not None and not sched_df.empty:
        sched_df.to_excel(xw, sheet_name="Programaciones", index=False)
    if runs_df is not None and not runs_df.empty:
        runs_df.to_excel(xw, sheet_name="Ejecuciones", index=False)

    hdr = PatternFill("solid", fgColor="1F4E78")
    for ws in xw.book.worksheets:
        hr = 4 if ws.title == "Resumen" else 1
        for c in ws[hr]:
            c.fill = hdr; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, col in enumerate(ws.columns, 1):
            w = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 55)
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(ws.max_column)}{ws.max_row}"
        if ws.title in ("Paridad ambientes",):
            for r in ws.iter_rows(min_row=2):
                for c in r[3:6]:
                    if c.value == "X": c.alignment = Alignment(horizontal="center"); c.fill = PatternFill("solid", fgColor="C6EFCE")

    if "Programaciones" in xw.book.sheetnames:
        ws_sched = xw.book["Programaciones"]
        col_estado = None
        for c in ws_sched[1]:
            if c.value == "Estado":
                col_estado = c.column
                break
        if col_estado:
            fills = {
                "habilitada": PatternFill("solid", fgColor="C6EFCE"),
                "deshabilitada": PatternFill("solid", fgColor="FFC7CE"),
                "vencida": PatternFill("solid", fgColor="FFEB9C"),
            }
            for r in ws_sched.iter_rows(min_row=2):
                val = r[col_estado - 1].value
                if val in fills:
                    r[col_estado - 1].fill = fills[val]

    if "Ejecuciones" in xw.book.sheetnames:
        ws_runs = xw.book["Ejecuciones"]
        col_status = None
        for c in ws_runs[1]:
            if c.value == "Estado":
                col_status = c.column
                break
        if col_status:
            fills = {
                "Completed": PatternFill("solid", fgColor="C6EFCE"),
                "Failed": PatternFill("solid", fgColor="FFC7CE"),
                "InProgress": PatternFill("solid", fgColor="FFEB9C"),
                "Cancelled": PatternFill("solid", fgColor="D9D9D9"),
            }
            for r in ws_runs.iter_rows(min_row=2):
                val = r[col_status - 1].value
                if val in fills:
                    r[col_status - 1].fill = fills[val]

    r = xw.book["Resumen"]
    r["A1"] = "Inventario Microsoft Fabric - Plan de transferencia"; r["A1"].font = Font(bold=True, size=14)
    parts = [f"Fuente: {src}", f"Total items: {len(d)}",
             f"Workspaces: {d.workspace.nunique()}", f"Tipos: {d.itemType.nunique()}"]
    if pipes_df is not None:
        parts.append(f"Pipelines: {len(pipes_df)}")
    if sched_df is not None:
        parts.append(f"Programaciones: {len(sched_df)}")
    if runs_df is not None:
        parts.append(f"Ejecuciones: {len(runs_df)}")
    r["A2"] = " | ".join(parts)

print("OK", out.resolve())
if sched_df is not None:
    print(f"  + Programaciones: {len(sched_df)} filas")
if runs_df is not None:
    print(f"  + Ejecuciones: {len(runs_df)} filas")
if pipes_df is not None:
    print(f"  + Pipelines: {len(pipes_df)} filas")
